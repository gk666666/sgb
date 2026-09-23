from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = ROOT / "defectlist.xlsx"


def load_mapping(path: Path) -> list[tuple[int, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows: list[tuple[int, str]] = []
    for index, row in enumerate(sheet.iter_rows(min_col=1, max_col=1, values_only=True), start=1):
        value = row[0]
        if value is None:
            continue
        code = str(value).strip()
        if code == "" or code.lower() in {"nan", "none"}:
            continue
        rows.append((index, code))
    return rows


def render_kusto_datatable(rows: list[tuple[int, str]]) -> str:
    lines = [
        ".set-or-replace dim_defectlist_5fn6 <|",
        "datatable(bit_pos:int, defect_code:string)",
        "[",
    ]
    for idx, code in rows:
        lines.append(f'    {idx}, "{code}",')
    if len(lines) > 3:
        lines[-1] = lines[-1].rstrip(",")
    lines.append("]")
    return "\n".join(lines)


def main():
    rows = load_mapping(DEFAULT_XLSX)
    print(f"# rows={len(rows)} source={DEFAULT_XLSX}")
    print(render_kusto_datatable(rows))


if __name__ == "__main__":
    main()
